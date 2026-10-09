from __future__ import annotations

import hashlib
import re
from dataclasses import dataclass
from datetime import date
from io import BytesIO
from typing import Any

from openpyxl import load_workbook

DATE_IN_NAME = re.compile(r"(?<!\d)(\d{1,2})[._-](\d{1,2})[._-](\d{2}|\d{4})(?!\d)")
PAIR_IN_SHEET = re.compile(r"^\s*Пара\s*(\d+)", re.IGNORECASE)
IGNORED_CELLS = {"ремонт", "свободно"}


@dataclass(frozen=True, slots=True)
class Lesson:
    lesson_date: date
    pair_number: int
    room: str
    group_name: str
    teacher: str

    @property
    def signature(self) -> str:
        value = "\0".join(
            (
                self.lesson_date.isoformat(),
                str(self.pair_number),
                self.room.casefold(),
                self.group_name.casefold(),
                self.teacher.casefold(),
            )
        )
        return hashlib.sha256(value.encode("utf-8")).hexdigest()


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return " ".join(str(value).split())


def _schedule_date(file_name: str) -> date:
    match = DATE_IN_NAME.search(file_name)
    if match is None:
        raise ValueError(f"Schedule filename has no date: {file_name}")
    day, month, year = (int(part) for part in match.groups())
    if year < 100:
        year += 2000
    return date(year, month, day)


def parse_schedule(file_name: str, content: bytes) -> list[Lesson]:
    """Read lessons from the two room/group/teacher blocks on each pair sheet."""
    lesson_date = _schedule_date(file_name)
    workbook = load_workbook(BytesIO(content), data_only=True, read_only=False)
    lessons: list[Lesson] = []

    try:
        for sheet_index, worksheet in enumerate(workbook.worksheets, start=1):
            match = PAIR_IN_SHEET.match(worksheet.title)
            pair_number = int(match.group(1)) if match else sheet_index

            for row_number, row in enumerate(
                worksheet.iter_rows(values_only=True), start=1
            ):
                # The first row contains the date, pair title, and occupancy note.
                if row_number == 1:
                    continue

                for block_start in (0, 3):
                    block = row[block_start : block_start + 3]
                    if len(block) < 3:
                        continue

                    room, group_name, teacher = (_cell_text(value) for value in block)
                    if not room or not group_name or not teacher:
                        continue
                    if group_name.casefold() in IGNORED_CELLS:
                        continue

                    lessons.append(
                        Lesson(
                            lesson_date=lesson_date,
                            pair_number=pair_number,
                            room=room,
                            group_name=group_name,
                            teacher=teacher,
                        )
                    )
    finally:
        workbook.close()

    return list(dict.fromkeys(lessons))
